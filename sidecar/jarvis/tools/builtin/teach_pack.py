"""Teach-by-showing Macros + Local File Search pack (Phase 12).

Two halves:

MACROS (teach by showing — at the TOOL-CALL level, not mouse wiggles)
    ``macro.record`` starts a capture session; the agent's subsequent tool
    calls are recorded (by the agent core, via the hook below) as a named
    macro. ``macro.stop`` ends the session and saves it to sqlite.
    ``macro.play`` replays the recorded tool-call sequence through the live
    agent's registry, so every replayed call passes the real policy and
    permission gates (high-risk replayed tools still ask for confirmation).
    Recorded args may contain ``{"var": "x"}`` placeholders that are filled
    from play-time ``args``.

    OUT OF SCOPE (deliberate): GUI-level recording — mouse moves, clicks,
    keystrokes — is NOT captured and NOT replayed. Screen automation is
    brittle, non-portable, and cannot be audited; recording tool calls is
    honest, deterministic, and replayable. The agent can always see the
    screen and choose to drive the GUI live instead.

    INTEGRATION NOTE FOR THE SIDECAR LOOP (parent: wire this in):
      1. In sidecar/jarvis/ipc/server.py, right after ``bind_agent(agent)``::

             from ..tools.builtin import teach_pack
             teach_pack.bind_agent(agent)

      2. In sidecar/jarvis/agent/core.py, in ``Agent.run``, right after::

             res = self.registry.call(name, args, audit=self.audit,
                                       confirmed=name in confirmed_tools)

         add::

             from ..tools.builtin import teach_pack
             observed = res.get("result", res.get("error", "awaiting confirmation"))
             teach_pack.note_call(name, args, observed)

         That is the ONLY capture hook: one call, no other core changes.
         Until the parent wires it, macro.record is honest about reporting
         that no session captured anything (sessions with zero steps are
         refused at stop time).

FILE SEARCH (local-only)
    ``files.index`` walks user document roots (pdf/docx/txt/md), extracts
    text, and builds a local sqlite FTS5 index. ``files.search`` runs ranked
    full-text queries ("find that March invoice") and returns path/snippet/
    rank hits. Everything stays on disk; needs_network=False for all tools
    in this pack. Extraction is dependency-gated: a missing pypdf or
    python-docx means that file type is SKIPPED and REPORTED, never faked.

SAFETY:
    - Handlers take dict -> return dict and never raise; failures are
      returned as {"error": "..."}.
    - macro.play is medium risk: replaying is real action, but each step is
      re-gated by the policy engine (a high-risk step inside a macro still
      requires confirmation instead of running blind).
    - Recording captures tool names, args, and short result summaries only.
      Recording skips any call whose tool name starts with "macro." so a
      macro can never capture itself (or another macro's control calls).
    - State lives in sqlite (stdlib only): ~/workspace/jarvis/data/teach.db,
      overridable with the JARVIS_TEACH_DB env var (tests use a tmp file).
      The recording session itself is in-memory and process-local.
"""
from __future__ import annotations

import json
import os
import sqlite3
import threading
from datetime import datetime, timezone
from pathlib import Path

HOME = Path.home()
JARVIS_DIR = HOME / "workspace" / "jarvis"
DEFAULT_DB_PATH = JARVIS_DIR / "data" / "teach.db"

NAME_RE = __import__("re").compile(r"^[a-z][a-z0-9_]{2,30}$")

# ---------------------------------------------------------------------------
# Agent binding (same pattern as builtin/workers_pack.py)
# ---------------------------------------------------------------------------
_AGENT = None


def bind_agent(agent) -> None:
    """Called once by the IPC server after the main Agent is constructed."""
    global _AGENT
    _AGENT = agent


# ---------------------------------------------------------------------------
# sqlite plumbing
# ---------------------------------------------------------------------------
_LOCK = threading.Lock()


def _db_path() -> Path:
    override = os.environ.get("JARVIS_TEACH_DB")
    return Path(override) if override else DEFAULT_DB_PATH


def _connect() -> sqlite3.Connection:
    path = _db_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path))
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS macros("
        " name TEXT PRIMARY KEY,"
        " created_at TEXT NOT NULL,"
        " steps_json TEXT NOT NULL)"
    )
    conn.execute(
        "CREATE VIRTUAL TABLE IF NOT EXISTS docs "
        "USING fts5(root, path UNINDEXED, content)"
    )
    return conn


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ---------------------------------------------------------------------------
# Recording hook (called by the agent core — parent wires the single call)
# ---------------------------------------------------------------------------
_RECORDING: dict | None = None  # {"name": str, "steps": [..], "started_at": str}
_RECORDING_LOCK = threading.Lock()


def recording_start(name: str) -> dict:
    """Begin capturing subsequent tool calls as macro ``name``."""
    with _RECORDING_LOCK:
        global _RECORDING
        if _RECORDING is not None:
            return {"error": (f"already recording macro '{_RECORDING['name']}'; "
                              "stop it with macro.stop first")}
        _RECORDING = {"name": name, "steps": [], "started_at": _now()}
        return {"recording": name, "started_at": _RECORDING["started_at"]}


def recording_stop() -> list[dict]:
    """End the capture session; return the captured steps (unsaved)."""
    with _RECORDING_LOCK:
        global _RECORDING
        if _RECORDING is None:
            return []
        steps = _RECORDING["steps"]
        _RECORDING = None
        return steps


def is_recording() -> bool:
    """True while a capture session is active."""
    with _RECORDING_LOCK:
        return _RECORDING is not None


def note_call(tool_name: str, args: dict, result_summary) -> None:
    """Record one tool call into the active session. No-op when idle.

    The agent core calls this after every registry.call. Calls to the macro
    control tools themselves ("macro.*") are never captured, so a macro
    cannot record itself.
    """
    with _RECORDING_LOCK:
        if _RECORDING is None:
            return
        if str(tool_name).startswith("macro."):
            return
        try:
            safe_args = json.loads(json.dumps(args or {}))
        except (TypeError, ValueError):
            safe_args = {"_unserializable_args": True}
        summary = str(result_summary) if result_summary is not None else ""
        _RECORDING["steps"].append({
            "tool": str(tool_name),
            "args": safe_args,
            "summary": summary[:500],
        })


# ---------------------------------------------------------------------------
# Macro handlers
# ---------------------------------------------------------------------------
def _validate_name(name) -> str | None:
    if not isinstance(name, str) or not NAME_RE.fullmatch(name):
        return ("name must match ^[a-z][a-z0-9_]{2,30}$ "
                "(lowercase, 3-30 chars, e.g. 'weekly_report')")
    return None


def macro_record_handler(args: dict) -> dict:
    """macro.record {name} — start a capture session."""
    try:
        name = args.get("name")
        err = _validate_name(name)
        if err:
            return {"error": err}
        out = recording_start(name)
        if "error" in out:
            return out
        return {"recording": name, "started_at": out["started_at"],
                "note": ("The agent core will now capture each tool call "
                         "(name, args, short result summary) into this macro. "
                         "Stop it with macro.stop.")}
    except Exception as exc:  # belt-and-braces: never raise
        return {"error": f"macro.record failed: {exc}"}


def macro_stop_handler(args: dict) -> dict:  # noqa: ARG001
    """macro.stop {} — end the session and save the macro to sqlite."""
    try:
        with _RECORDING_LOCK:
            sess = _RECORDING
        if sess is None:
            return {"recording": False,
                    "message": "no recording session is active"}
        name = sess["name"]
        steps = recording_stop()  # clears the session, returns steps
        if not steps:
            return {"recording": False,
                    "message": (f"session '{name}' captured no tool calls — "
                                "nothing saved (the agent core may not have "
                                "the note_call hook wired yet)")}
        with _LOCK:
            conn = _connect()
            try:
                conn.execute(
                    "INSERT OR REPLACE INTO macros(name, created_at, steps_json)"
                    " VALUES (?,?,?)",
                    (name, _now(), json.dumps(steps)))
                conn.commit()
            finally:
                conn.close()
        return {"macro": name, "steps": len(steps), "saved": True,
                "steps_detail": steps}
    except Exception as exc:
        return {"error": f"macro.stop failed: {exc}"}


def macro_list_handler(args: dict) -> dict:  # noqa: ARG001
    """macro.list {} — list saved macros with step counts."""
    try:
        with _LOCK:
            conn = _connect()
            try:
                rows = conn.execute(
                    "SELECT name, created_at, steps_json FROM macros ORDER BY name"
                ).fetchall()
            finally:
                conn.close()
        macros = []
        for name, created_at, steps_json in rows:
            try:
                steps = json.loads(steps_json)
            except (TypeError, ValueError):
                steps = []
            macros.append({"name": name, "steps": len(steps),
                           "created_at": created_at})
        return {"macros": macros, "count": len(macros),
                "recording_now": is_recording()}
    except Exception as exc:
        return {"error": f"macro.list failed: {exc}"}


_MISSING = object()


def _fill_vars(node, play_args: dict, missing: set) :
    """Recursively replace {"var": "x"} placeholders from play-time args."""
    if isinstance(node, dict):
        if set(node.keys()) == {"var"} and isinstance(node.get("var"), str):
            key = node["var"]
            if key in play_args:
                return play_args[key]
            missing.add(key)
            return _MISSING
        return {k: _fill_vars(v, play_args, missing) for k, v in node.items()}
    if isinstance(node, list):
        return [_fill_vars(v, play_args, missing) for v in node]
    return node


def _strip_missing(node):
    """Remove _MISSING sentinels left by unfilled vars (should not happen —
    unfilled vars are an error before replay)."""
    if node is _MISSING:
        return None
    if isinstance(node, dict):
        return {k: _strip_missing(v) for k, v in node.items()}
    if isinstance(node, list):
        return [_strip_missing(v) for v in node]
    return node


def macro_play_handler(args: dict) -> dict:
    """macro.play {name, args?} — replay a macro through the live registry.

    Each recorded call is executed via the bound agent's real registry, so
    policy and permission gates apply to every step: a high-risk step that
    needs confirmation will NOT run blind — the replay stops and reports
    it. Replay stops at the first error or confirmation gate.
    """
    try:
        name = args.get("name")
        err = _validate_name(name)
        if err:
            return {"error": err}
        play_args = args.get("args") or {}
        if not isinstance(play_args, dict):
            return {"error": "args must be an object of variable values"}

        with _LOCK:
            conn = _connect()
            try:
                row = conn.execute(
                    "SELECT steps_json FROM macros WHERE name=?", (name,)
                ).fetchone()
            finally:
                conn.close()
        if row is None:
            return {"error": f"no macro named '{name}' (see macro.list)"}
        try:
            steps = json.loads(row[0])
        except (TypeError, ValueError):
            return {"error": f"macro '{name}' has corrupt step data"}
        if not steps:
            return {"error": f"macro '{name}' has no steps"}

        if _AGENT is None:
            return {"error": "macro.play needs the live agent; not bound "
                             "(server.py must call teach_pack.bind_agent(agent))"}
        registry = getattr(_AGENT, "registry", None)
        if registry is None:
            return {"error": "macro.play needs the live agent; not bound "
                             "(bound agent exposes no .registry)"}

        results = []
        for i, step in enumerate(steps):
            missing: set = set()
            filled = _fill_vars(step.get("args", {}), play_args, missing)
            if missing:
                return {"error": (f"macro '{name}' step {i + 1} needs variables "
                                  f"{sorted(missing)} but they were not provided "
                                  f"in args"),
                        "macro": name, "completed_steps": results}
            filled = _strip_missing(filled)
            tool = step.get("tool", "?")
            try:
                res = registry.call(tool, filled, actor="agent")
            except Exception as exc:
                res = {"ok": False, "error": f"registry.call raised: {exc}"}
            step_out = {"step": i + 1, "tool": tool, "args": filled,
                        "result": res}
            results.append(step_out)
            if isinstance(res, dict) and (
                    res.get("error") or res.get("needs_confirmation")):
                return {"macro": name, "completed_steps": results,
                        "stopped_at_step": i + 1,
                        "reason": ("confirmation gate" if res.get("needs_confirmation")
                                   else "step error"),
                        "note": ("Replay stopped at the first gate/error — "
                                 "no further steps were executed.")}
        return {"macro": name, "completed_steps": results,
                "steps_executed": len(results)}
    except Exception as exc:
        return {"error": f"macro.play failed: {exc}"}


# ---------------------------------------------------------------------------
# File search handlers (local-only FTS5)
# ---------------------------------------------------------------------------
TEXT_EXTS = {".txt", ".md", ".markdown"}
MAX_FILE_BYTES = 50 * 1024 * 1024
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "env",
             ".tox", ".mypy_cache", ".pytest_cache"}


def _extract_text(path: Path) -> tuple[str | None, str | None]:
    """Return (text, skip_reason). text None + reason => skipped."""
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTS:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            return None, f"unreadable: {exc}"
    elif suffix == ".pdf":
        try:
            import pypdf  # noqa: PLC0415
        except ImportError:
            return None, "pypdf not installed — pdf skipped"
        try:
            reader = pypdf.PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as exc:
            return None, f"pdf extraction failed: {exc}"
    elif suffix == ".docx":
        try:
            import docx  # noqa: PLC0415
        except ImportError:
            return None, "python-docx not installed — docx skipped"
        try:
            doc = docx.Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            return None, f"docx extraction failed: {exc}"
    else:
        return None, f"unsupported extension '{suffix}'"
    if not text or not text.strip():
        return None, "no extractable text"
    return text, None


def files_index_handler(args: dict) -> dict:
    """files.index {roots} — walk doc roots, extract text, build FTS5 index.

    Respects excludes: hidden dirs, .git, node_modules, venvs, *.log,
    files >50MB. Re-indexing a root replaces its previous entries.
    """
    try:
        roots = args.get("roots")
        if not isinstance(roots, list) or not roots:
            return {"error": "roots must be a non-empty list of directory paths"}
        indexed = 0
        skipped: list[dict] = []
        done_roots: list[str] = []
        with _LOCK:
            conn = _connect()
            try:
                for raw in roots:
                    root = Path(raw).expanduser().resolve()
                    if not root.is_dir():
                        skipped.append({"path": str(raw),
                                        "reason": "not a directory"})
                        continue
                    root_key = str(root)
                    # FTS5: delete previous entries for this root via rowid.
                    conn.execute(
                        "DELETE FROM docs WHERE rowid IN "
                        "(SELECT rowid FROM docs WHERE root=?)", (root_key,))
                    for dirpath, dirnames, filenames in os.walk(root):
                        # Prune excluded / hidden dirs in place.
                        dirnames[:] = [d for d in dirnames
                                       if not d.startswith(".") and d not in SKIP_DIRS]
                        for fname in filenames:
                            fpath = Path(dirpath) / fname
                            if fpath.is_symlink():
                                continue
                            rel = fpath.relative_to(root)
                            if any(part.startswith(".") for part in rel.parts[:-1]):
                                continue
                            if fpath.suffix.lower() == ".log":
                                skipped.append({"path": str(fpath),
                                                "reason": "excluded: *.log"})
                                continue
                            if fpath.suffix.lower() not in TEXT_EXTS | {".pdf", ".docx"}:
                                continue  # unsupported types ignored silently
                            try:
                                size = fpath.stat().st_size
                            except OSError as exc:
                                skipped.append({"path": str(fpath),
                                                "reason": f"stat failed: {exc}"})
                                continue
                            if size > MAX_FILE_BYTES:
                                skipped.append({"path": str(fpath),
                                                "reason": f"excluded: >50MB ({size} bytes)"})
                                continue
                            text, reason = _extract_text(fpath)
                            if reason is not None:
                                skipped.append({"path": str(fpath), "reason": reason})
                                continue
                            conn.execute(
                                "INSERT INTO docs(root, path, content) VALUES (?,?,?)",
                                (root_key, str(fpath), text))
                            indexed += 1
                    conn.commit()
                    done_roots.append(root_key)
            finally:
                conn.close()
        return {"indexed": indexed, "skipped": skipped,
                "skipped_count": len(skipped), "roots": done_roots}
    except Exception as exc:
        return {"error": f"files.index failed: {exc}"}


def files_search_handler(args: dict) -> dict:
    """files.search {query, n?} — FTS5 ranked search over the local index."""
    try:
        query = (args.get("query") or "").strip()
        if not query:
            return {"error": "query must be a non-empty string"}
        n = args.get("n", 10)
        try:
            n = int(n)
        except (TypeError, ValueError):
            return {"error": "n must be an integer"}
        n = max(1, min(n, 50))
        # Quote each term individually and AND them: multi-word queries match
        # documents containing all terms in any order ("March invoice" finds
        # "Invoice March"), while quoting keeps FTS5 syntax chars harmless.
        terms = [t for t in query.split() if t]
        if not terms:
            return {"error": "query must be a non-empty string"}
        match_expr = "{content} : " + " AND ".join(
            '"' + t.replace('"', '""') + '"' for t in terms)
        with _LOCK:
            conn = _connect()
            try:
                try:
                    rows = conn.execute(
                        "SELECT path, snippet(docs, 2, '<<', '>>', '...', 40), "
                        "bm25(docs) AS rank FROM docs "
                        "WHERE docs MATCH ? ORDER BY rank LIMIT ?",
                        (match_expr, n)).fetchall()
                except sqlite3.OperationalError as exc:
                    return {"error": f"search query rejected by FTS5: {exc}",
                            "query": query}
            finally:
                conn.close()
        hits = [{"path": p, "snippet": s, "rank": r} for p, s, r in rows]
        return {"query": query, "hits": hits, "count": len(hits)}
    except Exception as exc:
        return {"error": f"files.search failed: {exc}"}


# ---------------------------------------------------------------------------
# Registration contract
# ---------------------------------------------------------------------------
TOOL_DEFS = [
    {"name": "macro.record",
     "description": ("Start recording a macro: the agent's NEXT tool calls are "
                     "captured (name, args, short result summary) as a named "
                     "macro. The agent core must call teach_pack.note_call() "
                     "after each registry call (parent wires one hook in "
                     "agent/core.py). Recording never captures macro.* control "
                     "calls, so a macro cannot record itself. Tool-call level "
                     "only — GUI (mouse/keyboard) recording is out of scope."),
     "handler": macro_record_handler, "risk": "low", "needs_network": False,
     "schema": {"name": "string"}},
    {"name": "macro.stop",
     "description": ("Stop the active recording session and save the captured "
                     "steps as a named macro in local sqlite. Sessions with "
                     "zero steps are not saved."),
     "handler": macro_stop_handler, "risk": "low", "needs_network": False,
     "schema": {}},
    {"name": "macro.play",
     "description": ("Replay a saved macro's recorded tool-call sequence. "
                     "Each step is executed through the LIVE agent's real "
                     "registry, so policy and permission gates still apply: "
                     "a high-risk step that needs confirmation will NOT run "
                     "blind — the replay stops and reports it. Replay stops "
                     "at the first error or gate. Recorded args may contain "
                     '{"var": "x"} placeholders filled from play-time args. '
                     "Requires teach_pack.bind_agent(agent) at startup."),
     "handler": macro_play_handler, "risk": "medium", "needs_network": False,
     "schema": {"name": "string", "args?": "object"}},
    {"name": "macro.list",
     "description": "List saved macros with step counts and creation time.",
     "handler": macro_list_handler, "risk": "low", "needs_network": False,
     "schema": {}},
    {"name": "files.index",
     "description": ("Walk document roots (pdf/docx/txt/md), extract text, and "
                     "build a LOCAL sqlite FTS5 index. Skips hidden dirs, "
                     ".git, node_modules, venvs, *.log, and files >50MB. "
                     "Extraction is dependency-gated: missing pypdf or "
                     "python-docx means that file type is skipped and "
                     "reported, never faked. Nothing leaves the machine."),
     "handler": files_index_handler, "risk": "low", "needs_network": False,
     "schema": {"roots": "array"}},
    {"name": "files.search",
     "description": ("Ranked full-text search over the local file index "
                     "('find that March invoice'). Returns path, snippet and "
                     "rank for each hit. Local only."),
     "handler": files_search_handler, "risk": "low", "needs_network": False,
     "schema": {"query": "string", "n?": "int"}},
]

# (risk, needs_network) additions for agent/policy.py's RISK_TABLE —
# without these, the policy engine default-denies unknown tools.
# Parent wiring:  from jarvis.agent.policy import RISK_TABLE
#                 RISK_TABLE.update(teach_pack.RISK_TABLE_ADDITIONS)
RISK_TABLE_ADDITIONS = {
    "macro.record": ("low", False),
    "macro.stop": ("low", False),
    "macro.play": ("medium", False),
    "macro.list": ("low", False),
    "files.index": ("low", False),
    "files.search": ("low", False),
}


def register(reg) -> None:
    """Wire the six Teach pack tools into a Registry."""
    from ..base import Tool
    for spec in TOOL_DEFS:
        reg.register(Tool(name=spec["name"], description=spec["description"],
                          schema=spec["schema"], handler=spec["handler"],
                          risk=spec["risk"], needs_network=spec["needs_network"]))
